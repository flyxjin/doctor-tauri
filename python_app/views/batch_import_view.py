from PyQt5.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
                             QFileDialog, QProgressBar, QTextEdit, QGroupBox, QMessageBox,
                             QTableWidget, QTableWidgetItem, QHeaderView)
from PyQt5.QtCore import Qt, QThread, pyqtSignal
from PyQt5.QtGui import QFont
import csv
import os

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class ImportWorker(QThread):
    progress = pyqtSignal(int, str)
    finished = pyqtSignal(int, int, int, list)
    
    def __init__(self, db, data_list):
        super().__init__()
        self.db = db
        self.data_list = data_list
        
    def run(self):
        added = 0
        updated = 0
        errors = []
        
        total = len(self.data_list)
        
        for i, data in enumerate(self.data_list):
            try:
                name = data.get('name', '').strip()
                if not name:
                    errors.append(f"第{i+1}行: 药材名称不能为空")
                    continue
                
                existing = self.db.fetchone(
                    "SELECT id FROM medicines WHERE name = ?", (name,)
                )
                
                if existing:
                    self.db.execute('''
                        UPDATE medicines 
                        SET alias=?, category=?, nature=?, taste=?, meridian=?,
                            efficacy=?, indications=?, usage=?, dosage=?, 
                            contraindication=?, notes=?
                        WHERE id=?
                    ''', (
                        data.get('alias', ''),
                        data.get('category', ''),
                        data.get('nature', ''),
                        data.get('taste', ''),
                        data.get('meridian', ''),
                        data.get('efficacy', ''),
                        data.get('indications', ''),
                        data.get('usage', ''),
                        data.get('dosage', ''),
                        data.get('contraindication', ''),
                        data.get('notes', ''),
                        existing[0]
                    ))
                    
                    if data.get('quantity') is not None:
                        self.db.execute('''
                            UPDATE inventory 
                            SET quantity=?, unit=?, price=?, min_stock=?, notes=?
                            WHERE medicine_id=?
                        ''', (
                            data.get('quantity', 0),
                            data.get('unit', 'g'),
                            data.get('price', 0),
                            data.get('min_stock', 10),
                            data.get('notes', ''),
                            existing[0]
                        ))
                    updated += 1
                else:
                    self.db.execute('''
                        INSERT INTO medicines 
                        (name, alias, category, nature, taste, meridian, 
                         efficacy, indications, usage, dosage, contraindication, notes)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (
                        name,
                        data.get('alias', ''),
                        data.get('category', ''),
                        data.get('nature', ''),
                        data.get('taste', ''),
                        data.get('meridian', ''),
                        data.get('efficacy', ''),
                        data.get('indications', ''),
                        data.get('usage', ''),
                        data.get('dosage', ''),
                        data.get('contraindication', ''),
                        data.get('notes', '')
                    ))
                    
                    med_id = self.db.fetchone(
                        "SELECT id FROM medicines WHERE name = ?", (name,)
                    )[0]
                    
                    self.db.execute('''
                        INSERT INTO inventory 
                        (medicine_id, quantity, unit, price, min_stock, notes)
                        VALUES (?, ?, ?, ?, ?, ?)
                    ''', (
                        med_id,
                        data.get('quantity', 0),
                        data.get('unit', 'g'),
                        data.get('price', 0),
                        data.get('min_stock', 10),
                        data.get('notes', '')
                    ))
                    added += 1
                    
            except Exception as e:
                errors.append(f"第{i+1}行: {str(e)}")
                
            self.progress.emit(int((i + 1) / total * 100), f"正在处理: {name}")
            
        self.finished.emit(added, updated, len(errors), errors)


class BatchImportView(QWidget):
    def __init__(self, db, parent=None):
        super().__init__(parent)
        self.db = db
        self.worker = None
        self.preview_data = []
        self.init_ui()
        
    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(15)
        
        title = QLabel('批量导入药材数据')
        title.setStyleSheet('font-size: 20px; font-weight: bold; color: #409eff;')
        layout.addWidget(title)
        
        file_group = QGroupBox('选择文件')
        file_layout = QHBoxLayout(file_group)
        
        self.file_label = QLabel('未选择文件')
        self.file_label.setStyleSheet('color: #909399;')
        
        self.select_btn = QPushButton('选择文件')
        self.select_btn.setStyleSheet('background-color: #409eff; color: white; padding: 8px 20px;')
        self.select_btn.clicked.connect(self.select_file)
        
        self.template_btn = QPushButton('下载模板')
        self.template_btn.setStyleSheet('background-color: #67c23a; color: white; padding: 8px 20px;')
        self.template_btn.clicked.connect(self.download_template)
        
        file_layout.addWidget(self.file_label, 1)
        file_layout.addWidget(self.select_btn)
        file_layout.addWidget(self.template_btn)
        layout.addWidget(file_group)
        
        self.preview_group = QGroupBox('数据预览')
        preview_layout = QVBoxLayout(self.preview_group)
        
        self.preview_table = QTableWidget()
        self.preview_table.setMaximumHeight(200)
        self.preview_table.setAlternatingRowColors(True)
        preview_layout.addWidget(self.preview_table)
        
        self.preview_info = QLabel()
        self.preview_info.setStyleSheet('color: #666;')
        preview_layout.addWidget(self.preview_info)
        
        self.preview_group.setVisible(False)
        layout.addWidget(self.preview_group)
        
        import_group = QGroupBox('导入操作')
        import_layout = QVBoxLayout(import_group)
        
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        
        self.status_label = QLabel('请选择要导入的文件')
        self.status_label.setStyleSheet('color: #666;')
        
        btn_layout = QHBoxLayout()
        self.import_btn = QPushButton('开始导入')
        self.import_btn.setStyleSheet('background-color: #67c23a; color: white; padding: 10px 30px;')
        self.import_btn.clicked.connect(self.start_import)
        self.import_btn.setEnabled(False)
        
        self.cancel_btn = QPushButton('取消')
        self.cancel_btn.setStyleSheet('background-color: #909399; color: white; padding: 10px 30px;')
        self.cancel_btn.clicked.connect(self.cancel_import)
        self.cancel_btn.setVisible(False)
        
        btn_layout.addStretch()
        btn_layout.addWidget(self.import_btn)
        btn_layout.addWidget(self.cancel_btn)
        btn_layout.addStretch()
        
        import_layout.addWidget(self.progress_bar)
        import_layout.addWidget(self.status_label)
        import_layout.addLayout(btn_layout)
        layout.addWidget(import_group)
        
        log_group = QGroupBox('导入日志')
        log_layout = QVBoxLayout(log_group)
        
        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        self.log_text.setMaximumHeight(150)
        log_layout.addWidget(self.log_text)
        
        layout.addWidget(log_group)
        
        help_group = QGroupBox('导入说明')
        help_layout = QVBoxLayout(help_group)
        help_text = QLabel('''
支持的文件格式: CSV、Excel (.xlsx)
必填字段: name (药材名称)
可选字段: alias (别名), category (分类), nature (药性), taste (药味), meridian (归经),
         efficacy (功效), indications (主治), usage (用法), dosage (用量),
         contraindication (禁忌), notes (备注), quantity (库存数量), unit (单位),
         price (单价), min_stock (最低库存)

注意事项:
1. 药材名称不能为空，重复名称将更新已有数据
2. CSV文件请使用UTF-8编码
3. 数值字段(库存、价格等)请填写数字
        ''')
        help_text.setStyleSheet('color: #666; line-height: 1.6;')
        help_layout.addWidget(help_text)
        layout.addWidget(help_group)
        
        layout.addStretch()
        
    def select_file(self):
        file_filter = '数据文件 (*.csv *.xlsx);;CSV文件 (*.csv);;Excel文件 (*.xlsx)'
        if not HAS_OPENPYXL:
            file_filter = 'CSV文件 (*.csv)'
            
        filename, _ = QFileDialog.getOpenFileName(self, '选择导入文件', '', file_filter)
        
        if not filename:
            return
            
        self.file_label.setText(filename)
        self.file_label.setStyleSheet('color: #333;')
        
        try:
            if filename.endswith('.csv'):
                self.preview_data = self.parse_csv(filename)
            elif filename.endswith('.xlsx') and HAS_OPENPYXL:
                self.preview_data = self.parse_excel(filename)
            else:
                QMessageBox.warning(self, '错误', '不支持的文件格式')
                return
                
            self.show_preview()
            self.import_btn.setEnabled(True)
            
        except Exception as e:
            QMessageBox.warning(self, '错误', f'文件解析失败: {str(e)}')
            
    def parse_csv(self, filename):
        data = []
        encodings = ['utf-8', 'utf-8-sig', 'gbk', 'gb2312']
        
        for encoding in encodings:
            try:
                with open(filename, 'r', encoding=encoding) as f:
                    reader = csv.DictReader(f)
                    headers = reader.fieldnames
                    
                    if not headers or 'name' not in [h.lower() for h in headers]:
                        raise ValueError('CSV文件必须包含name列')
                        
                    for row in reader:
                        normalized = {}
                        for key, value in row.items():
                            normalized[key.lower()] = value.strip() if value else ''
                        data.append(normalized)
                        
                break
            except UnicodeDecodeError:
                continue
                
        return data
        
    def parse_excel(self, filename):
        data = []
        wb = openpyxl.load_workbook(filename)
        ws = wb.active
        
        headers = [cell.value for cell in ws[1] if cell.value]
        headers_lower = [h.lower() for h in headers]
        
        if 'name' not in headers_lower:
            raise ValueError('Excel文件必须包含name列')
            
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            item = {}
            for i, header in enumerate(headers_lower):
                if i < len(row):
                    item[header] = str(row[i]) if row[i] else ''
            data.append(item)
            
        return data
        
    def show_preview(self):
        if not self.preview_data:
            return
            
        self.preview_group.setVisible(True)
        
        self.preview_table.clear()
        self.preview_table.setRowCount(min(5, len(self.preview_data)))
        
        if self.preview_data:
            headers = list(self.preview_data[0].keys())
            self.preview_table.setColumnCount(len(headers))
            self.preview_table.setHorizontalHeaderLabels(headers)
            
            for row_idx, row_data in enumerate(self.preview_data[:5]):
                for col_idx, header in enumerate(headers):
                    value = row_data.get(header, '')
                    item = QTableWidgetItem(str(value)[:30] if value else '')
                    self.preview_table.setItem(row_idx, col_idx, item)
                    
            self.preview_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
            
        self.preview_info.setText(f'共 {len(self.preview_data)} 条数据待导入')
        
    def download_template(self):
        filename, _ = QFileDialog.getSaveFileName(
            self, '保存模板文件', '药材导入模板.csv', 'CSV文件 (*.csv)'
        )
        
        if not filename:
            return
            
        headers = [
            'name', 'alias', 'category', 'nature', 'taste', 'meridian',
            'efficacy', 'indications', 'usage', 'dosage', 'contraindication',
            'notes', 'quantity', 'unit', 'price', 'min_stock'
        ]
        
        sample_data = [
            ['人参', '黄参', '补虚药', '温', '甘、微苦', '归脾、肺、心经',
             '大补元气', '体虚欲脱', '煎服', '3-9g', '实证忌服', '', '500', 'g', '85', '50'],
            ['黄芪', '黄耆', '补虚药', '微温', '甘', '归脾、肺经',
             '补气升阳', '气虚乏力', '煎服', '9-30g', '实证禁服', '', '600', 'g', '42', '60'],
        ]
        
        with open(filename, 'w', newline='', encoding='utf-8-sig') as f:
            writer = csv.writer(f)
            writer.writerow(headers)
            writer.writerows(sample_data)
            
        QMessageBox.information(self, '成功', f'模板已保存到:\n{filename}')
        
    def start_import(self):
        if not self.preview_data:
            return
            
        reply = QMessageBox.question(
            self, '确认导入',
            f'确定要导入 {len(self.preview_data)} 条数据吗？\n重复的药材名称将更新已有数据。',
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        )
        
        if reply != QMessageBox.Yes:
            return
            
        self.import_btn.setEnabled(False)
        self.select_btn.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.cancel_btn.setVisible(True)
        self.log_text.clear()
        
        self.worker = ImportWorker(self.db, self.preview_data)
        self.worker.progress.connect(self.on_progress)
        self.worker.finished.connect(self.on_finished)
        self.worker.start()
        
    def cancel_import(self):
        if self.worker and self.worker.isRunning():
            self.worker.terminate()
            self.log_text.append('导入已取消')
            
        self.reset_ui()
        
    def on_progress(self, percent, message):
        self.progress_bar.setValue(percent)
        self.status_label.setText(message)
        
    def on_finished(self, added, updated, errors, error_list):
        self.progress_bar.setValue(100)
        self.status_label.setText(f'导入完成！新增 {added} 条，更新 {updated} 条')
        
        self.log_text.append(f'导入完成！')
        self.log_text.append(f'  新增: {added} 条')
        self.log_text.append(f'  更新: {updated} 条')
        
        if errors > 0:
            self.log_text.append(f'  错误: {errors} 条')
            self.log_text.append('\n错误详情:')
            for err in error_list[:10]:
                self.log_text.append(f'  {err}')
            if len(error_list) > 10:
                self.log_text.append(f'  ... 还有 {len(error_list) - 10} 条错误')
                
        self.reset_ui()
        
    def reset_ui(self):
        self.import_btn.setEnabled(True)
        self.select_btn.setEnabled(True)
        self.cancel_btn.setVisible(False)
        self.worker = None