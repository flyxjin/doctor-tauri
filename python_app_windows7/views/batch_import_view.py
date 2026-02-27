# -*- coding: utf-8 -*-
"""
批量导入视图 - Windows 7兼容版本
支持CSV和Excel文件导入，增强数据验证和错误提示
"""
from PyQt5.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
                             QFileDialog, QProgressBar, QTextEdit, QGroupBox, QMessageBox,
                             QTableWidget, QTableWidgetItem, QHeaderView, QCheckBox)
from PyQt5.QtCore import Qt, QThread, pyqtSignal
from PyQt5.QtGui import QFont, QColor
import csv
import os
import logging

logger = logging.getLogger('MedicineSystem')

try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

from utils.excel_template import ExcelTemplateGenerator, DataValidator


def _get_value(data, key, default=''):
    if isinstance(data, dict):
        return data.get(key, default)
    elif hasattr(data, key):
        return getattr(data, key, default)
    return default


def _get_id_from_result(result):
    if result is None:
        return None
    if isinstance(result, dict):
        return result.get('id')
    elif len(result) > 0:
        return result[0]
    return None


class ImportWorker(QThread):
    progress = pyqtSignal(int, str)
    finished = pyqtSignal(int, int, int, list)
    
    def __init__(self, db, data_list):
        super().__init__()
        self.db = db
        self.data_list = data_list
        
    def run(self):
        added = 0
        updated = 0
        errors = []
        
        total = len(self.data_list)
        
        for i, data in enumerate(self.data_list):
            try:
                name = _get_value(data, 'name', '').strip()
                if not name:
                    errors.append(f"第{i+1}行: 药材名称不能为空")
                    continue
                
                existing = self.db.fetchone(
                    "SELECT id FROM medicines WHERE name = ?", (name,)
                )
                
                if existing:
                    existing_id = _get_id_from_result(existing)
                    self.db.execute('''
                        UPDATE medicines 
                        SET alias=?, category=?, nature=?, taste=?, meridian=?,
                            efficacy=?, indications=?, usage=?, dosage=?, 
                            contraindication=?, notes=?
                        WHERE id=?
                    ''', (
                        _get_value(data, 'alias', ''),
                        _get_value(data, 'category', ''),
                        _get_value(data, 'nature', ''),
                        _get_value(data, 'taste', ''),
                        _get_value(data, 'meridian', ''),
                        _get_value(data, 'efficacy', ''),
                        _get_value(data, 'indications', ''),
                        _get_value(data, 'usage', ''),
                        _get_value(data, 'dosage', ''),
                        _get_value(data, 'contraindication', ''),
                        _get_value(data, 'notes', ''),
                        existing_id
                    ))
                    
                    if _get_value(data, 'quantity') is not None:
                        self.db.execute('''
                            UPDATE inventory 
                            SET quantity=?, unit=?, price=?, min_stock=?, notes=?
                            WHERE medicine_id=?
                        ''', (
                            _get_value(data, 'quantity', 0),
                            _get_value(data, 'unit', 'g'),
                            _get_value(data, 'price', 0),
                            _get_value(data, 'min_stock', 10),
                            _get_value(data, 'notes', ''),
                            existing_id
                        ))
                    updated += 1
                else:
                    self.db.execute('''
                        INSERT INTO medicines 
                        (name, alias, category, nature, taste, meridian, 
                         efficacy, indications, usage, dosage, contraindication, notes)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    ''', (
                        name,
                        _get_value(data, 'alias', ''),
                        _get_value(data, 'category', ''),
                        _get_value(data, 'nature', ''),
                        _get_value(data, 'taste', ''),
                        _get_value(data, 'meridian', ''),
                        _get_value(data, 'efficacy', ''),
                        _get_value(data, 'indications', ''),
                        _get_value(data, 'usage', ''),
                        _get_value(data, 'dosage', ''),
                        _get_value(data, 'contraindication', ''),
                        _get_value(data, 'notes', '')
                    ))
                    
                    new_record = self.db.fetchone(
                        "SELECT id FROM medicines WHERE name = ?", (name,)
                    )
                    med_id = _get_id_from_result(new_record)
                    
                    self.db.execute('''
                        INSERT INTO inventory 
                        (medicine_id, quantity, unit, price, min_stock, notes)
                        VALUES (?, ?, ?, ?, ?, ?)
                    ''', (
                        med_id,
                        _get_value(data, 'quantity', 0),
                        _get_value(data, 'unit', 'g'),
                        _get_value(data, 'price', 0),
                        _get_value(data, 'min_stock', 10),
                        _get_value(data, 'notes', '')
                    ))
                    added += 1
                    
            except Exception as e:
                errors.append(f"第{i+1}行: {str(e)}")
                logger.error(f"导入数据失败: {e}")
                
            self.progress.emit(int((i + 1) / total * 100), f"正在处理: {name}")
            
        self.finished.emit(added, updated, len(errors), errors)


class BatchImportView(QWidget):
    def __init__(self, db, parent=None):
        super().__init__(parent)
        self.db = db
        self.worker = None
        self.preview_data = []
        self.validation_errors = []
        self.init_ui()
        
    def init_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(15)
        
        title = QLabel('批量导入药材数据')
        title.setStyleSheet('font-size: 20px; font-weight: bold; color: #409eff;')
        layout.addWidget(title)
        
        file_group = QGroupBox('选择文件')
        file_layout = QHBoxLayout(file_group)
        
        self.file_label = QLabel('未选择文件')
        self.file_label.setStyleSheet('color: #909399;')
        
        self.select_btn = QPushButton('选择文件')
        self.select_btn.setStyleSheet('background-color: #ffffff; color: #000000; border: 2px solid #000000; padding: 8px 20px;')
        self.select_btn.clicked.connect(self.select_file)
        
        self.template_btn = QPushButton('下载模板')
        self.template_btn.setStyleSheet('background-color: #ffffff; color: #000000; border: 1px solid #000000; padding: 8px 20px;')
        self.template_btn.clicked.connect(self.download_template)
        
        file_layout.addWidget(self.file_label, 1)
        file_layout.addWidget(self.select_btn)
        file_layout.addWidget(self.template_btn)
        layout.addWidget(file_group)
        
        self.preview_group = QGroupBox('数据预览')
        preview_layout = QVBoxLayout(self.preview_group)
        
        self.preview_table = QTableWidget()
        self.preview_table.setMaximumHeight(200)
        self.preview_table.setAlternatingRowColors(True)
        self.preview_table.verticalHeader().setVisible(False)
        preview_layout.addWidget(self.preview_table)
        
        preview_info_layout = QHBoxLayout()
        self.preview_info = QLabel()
        self.preview_info.setStyleSheet('color: #666;')
        preview_info_layout.addWidget(self.preview_info)
        
        self.error_info = QLabel()
        self.error_info.setStyleSheet('color: #f56c6c;')
        preview_info_layout.addWidget(self.error_info)
        preview_info_layout.addStretch()
        
        preview_layout.addLayout(preview_info_layout)
        
        self.preview_group.setVisible(False)
        layout.addWidget(self.preview_group)
        
        import_group = QGroupBox('导入操作')
        import_layout = QVBoxLayout(import_group)
        
        self.progress_bar = QProgressBar()
        self.progress_bar.setVisible(False)
        
        self.status_label = QLabel('请选择要导入的文件')
        self.status_label.setStyleSheet('color: #666;')
        
        btn_layout = QHBoxLayout()
        self.import_btn = QPushButton('开始导入')
        self.import_btn.setStyleSheet('background-color: #ffffff; color: #000000; border: 2px solid #000000; padding: 10px 30px;')
        self.import_btn.clicked.connect(self.start_import)
        self.import_btn.setEnabled(False)
        
        self.cancel_btn = QPushButton('取消')
        self.cancel_btn.setStyleSheet('background-color: #ffffff; color: #666666; border: 1px solid #e0e0e0; padding: 10px 30px;')
        self.cancel_btn.clicked.connect(self.cancel_import)
        self.cancel_btn.setVisible(False)
        
        btn_layout.addStretch()
        btn_layout.addWidget(self.import_btn)
        btn_layout.addWidget(self.cancel_btn)
        btn_layout.addStretch()
        
        import_layout.addWidget(self.progress_bar)
        import_layout.addWidget(self.status_label)
        import_layout.addLayout(btn_layout)
        layout.addWidget(import_group)
        
        log_group = QGroupBox('导入日志')
        log_layout = QVBoxLayout(log_group)
        
        self.log_text = QTextEdit()
        self.log_text.setReadOnly(True)
        self.log_text.setMaximumHeight(150)
        log_layout.addWidget(self.log_text)
        
        layout.addWidget(log_group)
        
        help_group = QGroupBox('导入说明')
        help_layout = QVBoxLayout(help_group)
        help_text = QLabel('''
<span style="color: #e74c3c; font-weight: bold;">必填字段:</span> name (药材名称) - 标红色表头
<span style="color: #409eff; font-weight: bold;">可选字段:</span> alias (别名), category (分类), nature (药性: 寒/热/温/凉/平), 
         taste (药味), meridian (归经), efficacy (功效), indications (主治), 
         usage (用法), dosage (用量), contraindication (禁忌), notes (备注), 
         quantity (库存数量), unit (单位), price (单价), min_stock (最低库存)

<span style="font-weight: bold;">支持的文件格式:</span> CSV、Excel (.xlsx)
<span style="font-weight: bold;">注意事项:</span>
1. 药材名称不能为空，重复名称将更新已有数据
2. CSV文件请使用UTF-8编码
3. 数值字段(库存、价格等)请填写数字，不能为负数
4. 药性字段请填写: 寒/热/温/凉/平
        ''')
        help_text.setStyleSheet('color: #666; line-height: 1.8;')
        help_text.setWordWrap(True)
        help_layout.addWidget(help_text)
        layout.addWidget(help_group)
        
        layout.addStretch()
        
    def select_file(self):
        file_filter = '数据文件 (*.csv *.xlsx);;CSV文件 (*.csv);;Excel文件 (*.xlsx)'
        if not HAS_OPENPYXL:
            file_filter = 'CSV文件 (*.csv)'
            
        filename, _ = QFileDialog.getOpenFileName(self, '选择导入文件', '', file_filter)
        
        if not filename:
            return
            
        self.file_label.setText(filename)
        self.file_label.setStyleSheet('color: #333;')
        
        try:
            if filename.endswith('.csv'):
                self.preview_data = self.parse_csv(filename)
            elif filename.endswith('.xlsx') and HAS_OPENPYXL:
                self.preview_data = self.parse_excel(filename)
            else:
                QMessageBox.warning(self, '错误', '不支持的文件格式')
                return
            
            valid_data, errors = DataValidator.validate_all(self.preview_data)
            self.validation_errors = errors
            
            self.show_preview()
            
            if errors:
                self.error_info.setText(f'发现 {len(errors)} 个验证错误')
                self.error_info.setStyleSheet('color: #f56c6c; font-weight: bold;')
                self.log_text.clear()
                self.log_text.append('<span style="color: #e74c3c;">数据验证错误:</span>')
                for err in errors[:20]:
                    self.log_text.append(f'  {err}')
                if len(errors) > 20:
                    self.log_text.append(f'  ... 还有 {len(errors) - 20} 个错误')
                
                reply = QMessageBox.warning(
                    self, '数据验证警告',
                    f'发现 {len(errors)} 个数据验证错误。\n\n'
                    f'有效数据: {len(valid_data)} 条\n'
                    f'错误数据: {len(self.preview_data) - len(valid_data)} 条\n\n'
                    f'是否继续导入有效数据？',
                    QMessageBox.Yes | QMessageBox.No,
                    QMessageBox.No
                )
                
                if reply == QMessageBox.Yes:
                    self.preview_data = valid_data
                    self.import_btn.setEnabled(len(valid_data) > 0)
                else:
                    self.import_btn.setEnabled(False)
                    return
            else:
                self.error_info.setText('')
                self.import_btn.setEnabled(True)
                
        except Exception as e:
            logger.error(f"文件解析失败: {e}")
            QMessageBox.warning(self, '错误', f'文件解析失败: {str(e)}')
            
    def parse_csv(self, filename):
        data = []
        encodings = ['utf-8', 'utf-8-sig', 'gbk', 'gb2312']
        
        for encoding in encodings:
            try:
                with open(filename, 'r', encoding=encoding) as f:
                    reader = csv.DictReader(f)
                    headers = reader.fieldnames
                    
                    if not headers or 'name' not in [h.lower() for h in headers]:
                        raise ValueError('CSV文件必须包含name列（药材名称）')
                        
                    for row in reader:
                        normalized = {}
                        for key, value in row.items():
                            normalized[key.lower()] = value.strip() if value else ''
                        data.append(normalized)
                        
                break
            except UnicodeDecodeError:
                continue
            except Exception as e:
                raise ValueError(f'CSV解析错误: {str(e)}')
                
        return data
        
    def parse_excel(self, filename):
        data = []
        wb = openpyxl.load_workbook(filename)
        ws = wb.active
        
        headers = [cell.value for cell in ws[1] if cell.value]
        headers_lower = [str(h).lower() for h in headers]
        
        if 'name' not in headers_lower:
            raise ValueError('Excel文件必须包含name列（药材名称）')
            
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row[0]:
                continue
            item = {}
            for i, header in enumerate(headers_lower):
                if i < len(row):
                    item[header] = str(row[i]) if row[i] is not None else ''
            data.append(item)
            
        return data
        
    def show_preview(self):
        if not self.preview_data:
            return
            
        self.preview_group.setVisible(True)
        
        self.preview_table.clear()
        self.preview_table.setRowCount(min(10, len(self.preview_data)))
        
        if self.preview_data:
            headers = list(self.preview_data[0].keys())
            self.preview_table.setColumnCount(len(headers))
            self.preview_table.setHorizontalHeaderLabels(headers)
            
            for row_idx, row_data in enumerate(self.preview_data[:10]):
                for col_idx, header in enumerate(headers):
                    value = row_data.get(header, '')
                    item = QTableWidgetItem(str(value)[:30] if value else '')
                    item.setTextAlignment(Qt.AlignCenter | Qt.AlignVCenter)
                    self.preview_table.setItem(row_idx, col_idx, item)
                    
            self.preview_table.horizontalHeader().setSectionResizeMode(QHeaderView.ResizeToContents)
            
        self.preview_info.setText(f'共 {len(self.preview_data)} 条数据待导入')
        
    def download_template(self):
        format_dialog = QMessageBox(self)
        format_dialog.setWindowTitle('选择模板格式')
        format_dialog.setText('请选择要下载的模板格式:')
        format_dialog.addButton('Excel格式 (.xlsx)', QMessageBox.AcceptRole)
        format_dialog.addButton('CSV格式 (.csv)', QMessageBox.RejectRole)
        format_dialog.addButton(QMessageBox.Cancel)
        
        result = format_dialog.exec_()
        
        if format_dialog.clickedButton().text() == 'Excel格式 (.xlsx)' and HAS_OPENPYXL:
            filename, _ = QFileDialog.getSaveFileName(
                self, '保存Excel模板', '药材导入模板.xlsx', 'Excel文件 (*.xlsx)'
            )
            if filename:
                if ExcelTemplateGenerator.generate_excel_template(filename):
                    QMessageBox.information(self, '成功', f'Excel模板已保存到:\n{filename}\n\n请参考模板格式填写数据。')
                else:
                    QMessageBox.warning(self, '失败', '生成Excel模板失败')
        elif format_dialog.clickedButton().text() == 'CSV格式 (.csv)' or not HAS_OPENPYXL:
            filename, _ = QFileDialog.getSaveFileName(
                self, '保存CSV模板', '药材导入模板.csv', 'CSV文件 (*.csv)'
            )
            if filename:
                if ExcelTemplateGenerator.generate_csv_template(filename):
                    QMessageBox.information(self, '成功', f'CSV模板已保存到:\n{filename}\n\n请参考模板格式填写数据。')
                else:
                    QMessageBox.warning(self, '失败', '生成CSV模板失败')
        
    def start_import(self):
        if not self.preview_data:
            return
            
        reply = QMessageBox.question(
            self, '确认导入',
            f'确定要导入 {len(self.preview_data)} 条数据吗？\n\n'
            f'<span style="color: #e74c3c;">重复的药材名称将更新已有数据。</span>',
            QMessageBox.Yes | QMessageBox.No, QMessageBox.No
        )
        
        if reply != QMessageBox.Yes:
            return
            
        self.import_btn.setEnabled(False)
        self.select_btn.setEnabled(False)
        self.progress_bar.setVisible(True)
        self.cancel_btn.setVisible(True)
        self.log_text.clear()
        
        self.worker = ImportWorker(self.db, self.preview_data)
        self.worker.progress.connect(self.on_progress)
        self.worker.finished.connect(self.on_finished)
        self.worker.start()
        
    def cancel_import(self):
        if self.worker and self.worker.isRunning():
            self.worker.terminate()
            self.log_text.append('<span style="color: #e74c3c;">导入已取消</span>')
            
        self.reset_ui()
        
    def on_progress(self, percent, message):
        self.progress_bar.setValue(percent)
        self.status_label.setText(message)
        
    def on_finished(self, added, updated, errors, error_list):
        self.progress_bar.setValue(100)
        self.status_label.setText(f'导入完成！新增 {added} 条，更新 {updated} 条')
        
        self.log_text.append('<span style="color: #67c23a; font-weight: bold;">导入完成！</span>')
        self.log_text.append(f'  <span style="color: #67c23a;">新增: {added} 条</span>')
        self.log_text.append(f'  <span style="color: #409eff;">更新: {updated} 条</span>')
        
        if errors > 0:
            self.log_text.append(f'  <span style="color: #e74c3c;">错误: {errors} 条</span>')
            self.log_text.append('\n<span style="font-weight: bold;">错误详情:</span>')
            for err in error_list[:10]:
                self.log_text.append(f'  {err}')
            if len(error_list) > 10:
                self.log_text.append(f'  ... 还有 {len(error_list) - 10} 条错误')
        
        logger.info(f"批量导入完成: 新增{added}条, 更新{updated}条, 错误{errors}条")
                
        self.reset_ui()
        
    def reset_ui(self):
        self.import_btn.setEnabled(True)
        self.select_btn.setEnabled(True)
        self.cancel_btn.setVisible(False)
        self.worker = None
